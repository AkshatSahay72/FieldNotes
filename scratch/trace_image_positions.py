import docx
import re

doc = docx.Document('d:/Coding/Major Project/FieldNotes/FieldNotes_MongoDB_Project_Report.docx')

for i, tbl in enumerate(doc.tables):
    xml = tbl._tbl.xml
    embeds = re.findall(r'r:embed="([^"]+)"', xml)
    if embeds:
        # Find preceding paragraphs
        # Search backwards in doc elements
        print(f"\n==========================================")
        print(f"Table {i+1} has images: {embeds}")
        # Let's find where this table is in doc.element.body
        tbl_elm = tbl._tbl
        prev_elm = tbl_elm.getprevious()
        pre_texts = []
        while prev_elm is not None and len(pre_texts) < 4:
            if prev_elm.tag.endswith('p'):
                txt = prev_elm.text.strip() if prev_elm.text else ""
                # extract all w:t inside prev_elm
                ts = re.findall(r'<w:t[^>]*>(.*?)</w:t>', prev_elm.xml)
                if ts:
                    pre_texts.append("".join(ts))
            prev_elm = prev_elm.getprevious()
        print("Preceding elements:")
        for pt in reversed(pre_texts):
            print(f"  [P]: {pt[:120]}")
        
        # Check if table has text or caption
        cell_txt = tbl.cell(0,0).text.strip()
        print(f"Cell text: {repr(cell_txt)}")
        
        next_elm = tbl_elm.getnext()
        post_texts = []
        while next_elm is not None and len(post_texts) < 4:
            if next_elm.tag.endswith('p'):
                ts = re.findall(r'<w:t[^>]*>(.*?)</w:t>', next_elm.xml)
                if ts:
                    post_texts.append("".join(ts))
            next_elm = next_elm.getnext()
        print("Following elements:")
        for pt in post_texts:
            print(f"  [P]: {pt[:120]}")
