import docx

doc = docx.Document('d:/Coding/Major Project/FieldNotes/FieldNotes_MongoDB_Project_Report.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Check Section 8 headings and text
in_sec_8 = False
for p in doc.paragraphs:
    t = p.text.strip()
    if "8. RESULTS AND OUTPUT ANALYSIS" in t:
        in_sec_8 = True
    if "9. DISCUSSION" in t:
        in_sec_8 = False
    if in_sec_8:
        if t.startswith("8.") or "Figure" in t or "Table" in t or "Screenshot" in t.lower() or "output" in t.lower() or "objective" in t.lower():
            print(f"  P: {t[:80]}")
