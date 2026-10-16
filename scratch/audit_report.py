import docx
import re

doc_path = 'd:/Coding/Major Project/FieldNotes/FieldNotes_MongoDB_Project_Report.docx'
doc = docx.Document(doc_path)

print("=== 1. PRELIMINARIES PARAGRAPHS (FIRST 140) ===")
for i in range(min(140, len(doc.paragraphs))):
    txt = doc.paragraphs[i].text.strip()
    if any(k in txt for k in ["TABLE OF CONTENTS", "LIST OF FIGURES", "LIST OF TABLES", "1. INTRODUCTION"]):
        print(f"P[{i}]: {txt}")

# Check all figure references in tables
print("\n=== 2. TABLES INSPECTION ===")
for idx, tbl in enumerate(doc.tables):
    for r_idx, r in enumerate(tbl.rows):
        for c_idx, c in enumerate(r.cells):
            txt = c.text.strip()
            if "Figure " in txt or "PLACEHOLDER" in txt:
                print(f"Tbl[{idx+1}] R{r_idx}C{c_idx}: {txt[:80]}...")

# Codebase terms
print("\n=== 3. CODEBASE CONSISTENCY AUDIT ===")
all_text = []
for p in doc.paragraphs:
    all_text.append(p.text)
for tbl in doc.tables:
    for r in tbl.rows:
        for c in r.cells:
            all_text.append(c.text)
full_body = "\n".join(all_text)

terms_to_check = [
    "fieldnote_db", "fieldnote_app",
    "idx_notes_vector", "vector_index",
    "all-MiniLM-L6-v2", "MiniLM",
    "384", "1536", "768",
    "fs.files", "fs.chunks",
    "HTTP 422", "HTTP 400",
    "single-owner", "owner_id", "user_id",
    "bcrypt", "Argon2",
    "JWT", "session",
    "Render", "Atlas"
]

print("Occurrence counts for key terms:")
for term in terms_to_check:
    cnt = len(re.findall(r'\b' + re.escape(term) + r'\b', full_body, re.IGNORECASE))
    print(f"  {term}: {cnt}")
