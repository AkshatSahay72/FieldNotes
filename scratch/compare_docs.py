import docx
import re

cur_doc = docx.Document('d:/Coding/Major Project/FieldNotes/FieldNotes_MongoDB_Project_Report.docx')
new_doc = docx.Document('scratch/test_report.docx')

print(f"Current doc: {len(cur_doc.paragraphs)} paragraphs, {len(cur_doc.tables)} tables")
print(f"Fresh build: {len(new_doc.paragraphs)} paragraphs, {len(new_doc.tables)} tables")

# Check all figure references in fresh build
print("\n--- Fresh build figures ---")
for idx, tbl in enumerate(new_doc.tables):
    for r in tbl.rows:
        for c in r.cells:
            for p in c.paragraphs:
                txt = p.text
                if "Figure " in txt:
                    m = re.search(r'Figure\s+(\d+\.\d+)[:\s]+([^\n]+)', txt)
                    if m:
                        print(f"  Table[{idx+1}]: Fig {m.group(1)} - {m.group(2)[:60]}")
